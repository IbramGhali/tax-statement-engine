"""Reviewer workbook. Summary cells are formulas over the Transactions sheet, so a reviewer
can trace every figure to the rows behind it instead of trusting a pasted number."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from tax_engine.calc import Reconciliation
from tax_engine.models import FIGURE_LABELS, Figures, Finding, Snapshot

TX_HEADERS = ["id", "date", "kind", "tax_code", "net", "tax", "included", "reason"]
MONEY_FORMAT = "#,##0.00"

# Column letters on the Transactions sheet.
KIND, CODE, NET, TAX, INCLUDED = "C", "D", "E", "F", "G"


def _sumifs(column: str, **criteria: str) -> str:
    cols = {"kind": KIND, "code": CODE, "included": INCLUDED}
    parts = [f"Transactions!{column}:{column}"]
    for name, value in criteria.items():
        parts.append(f'Transactions!{cols[name]}:{cols[name]},"{value}"')
    return "SUMIFS(" + ",".join(parts) + ")"


SUMMARY_FORMULAS: dict[str, str] = {
    "sales_total": f"={_sumifs(NET, kind='sale', included='yes')}"
    f"+{_sumifs(TAX, kind='sale', included='yes')}",
    "sales_taxable": f"={_sumifs(NET, kind='sale', code='STD', included='yes')}",
    "sales_exempt": f"={_sumifs(NET, kind='sale', code='EXEMPT', included='yes')}",
    "output_tax": f"={_sumifs(TAX, kind='sale', code='STD', included='yes')}",
    "purchases_total": f"={_sumifs(NET, kind='purchase', included='yes')}"
    f"+{_sumifs(TAX, kind='purchase', included='yes')}",
    "input_tax": f"={_sumifs(TAX, kind='purchase', code='STD', included='yes')}",
    "wages_gross": "=Payroll!B2",
    "wages_withheld": "=Payroll!B3",
    "net_payable": "=B5-B7+B9",
}
SUMMARY_ROWS = list(SUMMARY_FORMULAS)  # row 2 onwards, in this order


def build_workbook(
    path: Path,
    business: str,
    snapshot: Snapshot,
    rec: Reconciliation,
    figures: Figures,
    findings: list[Finding],
    status: str,
    commentary: str,
    versions: dict[str, str],
) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Line", "Amount", "Engine value (check)"])
    for name in SUMMARY_ROWS:
        summary.append([FIGURE_LABELS[name], SUMMARY_FORMULAS[name], float(getattr(figures, name))])
    summary.append([])
    summary.append(["Business", business])
    summary.append(["Quarter", snapshot.quarter])
    summary.append(["Status", status])
    summary.append(["Reviewer note", commentary])
    for row in summary.iter_rows(min_row=2, max_row=1 + len(SUMMARY_ROWS), min_col=2, max_col=3):
        for cell in row:
            cell.number_format = MONEY_FORMAT

    tx = wb.create_sheet("Transactions")
    tx.append(TX_HEADERS)
    excluded = {t.id: "uncoded" for t in rec.uncoded} | {
        t.id: "outside period" for t in rec.out_of_period
    }
    for t in sorted(snapshot.transactions, key=lambda t: (t.date, t.id)):
        tx.append(
            [
                t.id,
                t.date.isoformat(),
                t.kind,
                t.tax_code or "",
                float(t.net),
                float(t.tax),
                "no" if t.id in excluded else "yes",
                excluded.get(t.id, ""),
            ]
        )

    payroll = wb.create_sheet("Payroll")
    payroll.append(["Item", "Amount"])
    payroll.append(["Gross wages", float(snapshot.payroll.gross_wages)])
    payroll.append(["Withheld", float(snapshot.payroll.withheld)])

    fsheet = wb.create_sheet("Findings")
    fsheet.append(["Rule", "Severity", "Message"])
    for f in findings:
        fsheet.append([f.rule_id, f.severity.value, f.message])

    meta = wb.create_sheet("Provenance")
    for key, value in versions.items():
        meta.append([key, value])

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for i, column in enumerate(ws.columns, start=1):
            width = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
